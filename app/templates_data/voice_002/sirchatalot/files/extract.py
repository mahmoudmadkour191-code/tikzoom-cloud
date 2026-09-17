'''
Extract plain text from uploaded documents (PDF, DOCX, DOC, PPTX, PPT, TXT-like).
Legacy .doc/.ppt go through catdoc/catppt (installed in the Docker image).
'''

import asyncio

import aiofiles
import docx
import fitz  # PyMuPDF
import pptx

from sirchatalot.logging_setup import get_logger

logger = get_logger('extract')

SUPPORTED_EXTENSIONS = ('.pdf', '.docx', '.doc', '.pptx', '.ppt', '.txt', '.md', '.csv', '.log')


class FilesProcessor:
    async def convert_to_text(self, file_path: str) -> str | None:
        try:
            lower = file_path.lower()
            if lower.endswith('.pdf'):
                return await self._extract_pdf(file_path)
            if lower.endswith('.docx'):
                return await asyncio.to_thread(self._extract_docx, file_path)
            if lower.endswith('.pptx'):
                return await asyncio.to_thread(self._extract_pptx, file_path)
            if lower.endswith('.doc'):
                return await self._run_converter('catdoc', file_path)
            if lower.endswith('.ppt'):
                return await self._run_converter('catppt', file_path)
            if lower.endswith(('.txt', '.md', '.csv', '.log')):
                async with aiofiles.open(file_path, mode='r', encoding='utf-8') as f:
                    return await f.read()
            logger.warning(f'Unsupported file type: {file_path}')
            return None
        except Exception:
            logger.exception(f'Error converting file to text: {file_path}')
            return None

    async def _extract_pdf(self, file_path: str) -> str:
        doc = await asyncio.to_thread(fitz.open, file_path)
        try:
            parts = []
            for page_num in range(len(doc)):
                parts.append(await asyncio.to_thread(doc[page_num].get_text))
            return '\n'.join(parts)
        finally:
            await asyncio.to_thread(doc.close)

    def _extract_docx(self, file_path: str) -> str:
        document = docx.Document(file_path)
        return '\n'.join(para.text for para in document.paragraphs)

    def _extract_pptx(self, file_path: str) -> str:
        presentation = pptx.Presentation(file_path)
        texts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text:
                    texts.append(shape.text)
        return '\n'.join(texts)

    async def _run_converter(self, binary: str, file_path: str) -> str | None:
        try:
            proc = await asyncio.create_subprocess_exec(
                binary, file_path,
                stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
            )
            stdout, stderr = await proc.communicate()
            if stderr:
                logger.warning(f'{binary} warning for {file_path}: {stderr.decode(errors="replace")}')
            return stdout.decode(errors='replace')
        except FileNotFoundError:
            logger.error(f'{binary} is not installed; cannot convert {file_path}')
            return None
